"""One-time script to create assets/templates/laporan_bulanan_template.xlsx
from the reference Laporan Bulanan April.xlsx.

Keeps rows 1-2 (headers) and rows 3-4 (one sample data row + one sample
Total Personal row) — these serve as style references for the generator.
Deletes all rows from row 5 onwards.

Run from project root:
    python tools/build_template.py [path/to/reference.xlsx]

If no path is given, defaults to 'Data Absensi/Laporan Bulanan April.xlsx'.
"""
import sys
from pathlib import Path
from openpyxl import load_workbook


def main():
    if len(sys.argv) > 1:
        ref = Path(sys.argv[1])
    else:
        ref = Path("Data Absensi/Laporan Bulanan April.xlsx")
    if not ref.exists():
        print(f"ERROR: Reference file not found: {ref}", file=sys.stderr)
        sys.exit(1)

    out = Path("assets/templates/laporan_bulanan_template.xlsx")
    out.parent.mkdir(exist_ok=True, parents=True)

    wb = load_workbook(ref)
    ws = wb.active

    # We keep rows 1, 2, 3, 4. Delete row 5 onwards.
    if ws.max_row > 4:
        ws.delete_rows(5, ws.max_row - 4)

    wb.save(out)
    print(f"Wrote {out} (max_row now: {ws.max_row})")


if __name__ == "__main__":
    main()

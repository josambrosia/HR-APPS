"""Analisa tambahan: kolom Lembur + kandidat lembur September klien."""
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.path.insert(0, r"D:\Gawe\Project X\HR App")
from openpyxl import load_workbook

SRC = Path(r"D:\Gawe\Project X\HR App\Export\Laporan Bulanan September 2026 [Auto Filled].xlsx")
ws = load_workbook(SRC, data_only=True).active

rows = []
for r in range(3, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    if not a or "Total" in str(a):
        continue
    tgl = ws.cell(row=r, column=3).value
    if isinstance(tgl, datetime):
        tgl = tgl.date().isoformat()
    rows.append({
        "nama": a, "tanggal": str(tgl),
        "masuk": ws.cell(row=r, column=7).value,
        "keluar": ws.cell(row=r, column=8).value,
        "lembur": ws.cell(row=r, column=10).value,
        "terlambat": ws.cell(row=r, column=12).value,
    })

print("kolom Lembur (dari file fingerprint) — distribusi:")
for v, n in Counter(str(r["lembur"]) for r in rows).most_common(8):
    print(f"  {v!r}: {n}")

sep = [r for r in rows if r["tanggal"] >= "2026-09-14"]
cand = [r for r in sep if r["keluar"] and str(r["keluar"]) > "18:00"]
print(f"\nkandidat lembur (keluar > 18:00) >= 14 Sep: {len(cand)}")
for r in cand[:20]:
    print(f"  {r['nama'][:16]:16} {r['tanggal']} keluar={r['keluar']} "
          f"lembur_file={r['lembur']!r}")

# Lembur file vs hitungan app (keluar - 17:00)
def to_min(v):
    if v is None:
        return None
    try:
        h, m = str(v).replace(".", ":").split(":")
        return int(h) * 60 + int(m)
    except (ValueError, AttributeError):
        return None

print("\nbandingkan 'Lembur' file vs hitungan app (keluar-17:00) "
      "untuk kandidat:")
for r in cand[:12]:
    km = to_min(r["keluar"])
    app = (km - 17 * 60) if km else None
    print(f"  {r['nama'][:16]:16} {r['tanggal']} keluar={r['keluar']} "
          f"file={r['lembur']!r} app={app} mnt")

before = [r for r in rows if r["tanggal"] < "2026-09-14"]
print(f"\nsebelum 14 Sep: {len(before)} baris, "
      f"kandidat keluar>18:00: "
      f"{sum(1 for r in before if r['keluar'] and str(r['keluar']) > '18:00')}")

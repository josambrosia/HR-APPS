"""Bedah Laporan Bulanan September 2026 milik klien (output aplikasi)."""
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.path.insert(0, r"D:\Gawe\Project X\HR App")

from openpyxl import load_workbook

SRC = Path(r"D:\Gawe\Project X\HR App\Export\Laporan Bulanan September 2026 [Auto Filled].xlsx")
wb = load_workbook(SRC, data_only=True)
ws = wb.active
print("sheet:", ws.title, "dims:", ws.dimensions)

# Header: lihat 2 baris pertama
for r in range(1, 3):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 18)]
    print(f"row{r}:", vals)

# Kolom: A Nama, B Dept, C Tanggal, D Hari, E Tipe, F Jadwal, G Masuk,
# H Keluar, I Kerja, J Lembur, K Kurang, L Terlambat, M PulangCepat,
# N Absen, O Lupa, P Ijin, Q Alasan

rows = []
for r in range(3, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    if not a or "Total" in str(a):
        continue
    tgl = ws.cell(row=r, column=3).value
    if isinstance(tgl, datetime):
        tgl = tgl.date().isoformat()
    elif tgl is not None:
        tgl = str(tgl)
    rows.append({
        "nama": a, "tanggal": tgl,
        "tipe": ws.cell(row=r, column=5).value,
        "jadwal": ws.cell(row=r, column=6).value,
        "masuk": ws.cell(row=r, column=7).value,
        "keluar": ws.cell(row=r, column=8).value,
        "lembur": ws.cell(row=r, column=10).value,
        "terlambat": ws.cell(row=r, column=12).value,
        "alasan": ws.cell(row=r, column=17).value,
    })

print(f"\ntotal baris data: {len(rows)}")

# Nilai kolom Jadwal (dari file fingerprint — menunjukkan jadwal MESIN)
print("\nkolom 'Jadwal' (dari file fingerprint):")
for v, n in Counter(str(r["jadwal"]) for r in rows).most_common(10):
    print(f"  {v!r}: {n}")

print("\nkolom Tipe:")
for v, n in Counter(str(r["tipe"]) for r in rows).most_common(10):
    print(f"  {v!r}: {n}")

# Fokus tanggal >= 2026-09-14
sep = [r for r in rows if r["tanggal"] and r["tanggal"] >= "2026-09-14"]
print(f"\nbaris >= 14 Sep: {len(sep)}")

print("\ncontoh baris >= 14 Sep (Nama, tgl, jadwal, masuk, keluar, terlambat):")
for r in sep[:25]:
    print(f"  {r['nama'][:18]:18} {r['tanggal']} jadwal={r['jadwal']!r} "
          f"masuk={r['masuk']!r} keluar={r['keluar']!r} "
          f"telat={r['terlambat']!r}")

# Analisa: apakah 'Terlambat' cocok dengan hitungan 08:00 atau 09:00?
def to_min(v):
    if v is None:
        return None
    s = str(v).strip()
    try:
        h, m = s.replace(".", ":").split(":")
        return int(h) * 60 + int(m)
    except (ValueError, AttributeError):
        return None

agree_08 = agree_09 = other = unknown = 0
examples = []
for r in sep:
    ms = to_min(r["masuk"])
    tl = r["terlambat"]
    if ms is None or not isinstance(tl, (int, float)):
        unknown += 1
        continue
    d08 = max(0, ms - 8 * 60)
    d09 = max(0, ms - 9 * 60)
    if abs(tl - d08) <= 1:
        agree_08 += 1
        if len(examples) < 12:
            examples.append((r["masuk"], int(tl), "08.00-based"))
    elif abs(tl - d09) <= 1:
        agree_09 += 1
        if len(examples) < 12:
            examples.append((r["masuk"], int(tl), "09.00-based"))
    else:
        other += 1
        if len(examples) < 12:
            examples.append((r["masuk"], int(tl), "TIDAK COCOK"))

print(f"\nAnalisa 'Terlambat' vs jam masuk (>= 14 Sep):")
print(f"  cocok hitungan 08.00 : {agree_08}")
print(f"  cocok hitungan 09.00 : {agree_09}")
print(f"  tidak cocok          : {other}")
print(f"  kosong/tidak lengkap : {unknown}")
print("\ncontoh (masuk, terlambat, kesimpulan):")
for e in examples:
    print("  ", e)

# Masuk jam berapa saja >= 14 Sep?
print("\ndistribusi jam Masuk >= 14 Sep (10 terbanyak):")
for v, n in Counter(str(r["masuk"]) for r in sep).most_common(10):
    print(f"  {v!r}: {n}")

# Apakah ada baris yang seharusnya TIDAK telat (masuk <= 09:12) tapi
# tetap diekspor telat?
late_but_ok = [r for r in sep
               if to_min(r["masuk"]) is not None
               and to_min(r["masuk"]) <= 9 * 60 + 12
               and isinstance(r["terlambat"], (int, float))
               and r["terlambat"] > 0]
print(f"\nbaris masuk <= 09:12 TAPI terlanjut telat>0: {len(late_but_ok)}")
for r in late_but_ok[:15]:
    print(f"  {r['nama'][:18]:18} {r['tanggal']} masuk={r['masuk']} "
          f"telat={r['terlambat']} alasan={r['alasan']!r}")

# Bandingkan periode sebelum cutoff
before = [r for r in rows if r["tanggal"] and r["tanggal"] < "2026-09-14"]
print(f"\nbaris < 14 Sep: {len(before)}")
print("contoh masuk < 14 Sep:", [r['masuk'] for r in before[:12]])

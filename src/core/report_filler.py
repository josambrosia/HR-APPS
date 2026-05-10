import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Tuple

from openpyxl import load_workbook

from src.core.reason_mapper import render_alasan_ijin
from src.parsers.helpers import parse_date_id

ALASAN_IJIN_COL = 17  # Q


@dataclass
class FillSummary:
    filled_count: int
    na_count: int
    not_found_count: int


def fill_monthly_report(
    xlsx_path: Path, conn: sqlite3.Connection
) -> Tuple[Path, FillSummary]:
    """Open Laporan Bulanan, write Alasan Ijin column from DB, save as <stem> [filled].xlsx.

    - Original file is NOT touched.
    - Only column Q is modified; other cells & formatting preserved.
    - Rows where employee+date not found in DB → leave as is (count not_found).
    - Rows where DB has has_issue=1 but no reason → write 'NA / Belum ada kabar'.
    - Rows with reason → render via reason_mapper.
    """
    wb = load_workbook(xlsx_path)
    ws = wb.active

    summary = FillSummary(0, 0, 0)

    # rows 1 = header, 2 = sub-unit. Data starts at row 3.
    for r in range(3, ws.max_row + 1):
        nama_cell = ws.cell(row=r, column=1).value
        tgl_cell = ws.cell(row=r, column=3).value
        if not nama_cell:
            continue
        nama = str(nama_cell).strip()
        tanggal = parse_date_id(tgl_cell)
        if not tanggal:
            continue

        row = conn.execute(
            """
            SELECT ar.has_issue, ar.reason_category, ar.reason_detail
              FROM attendance_records ar
              JOIN employees e ON ar.employee_id = e.id
             WHERE e.nama = ? COLLATE NOCASE AND ar.tanggal = ?
            """,
            (nama, tanggal),
        ).fetchone()

        if row is None:
            summary.not_found_count += 1
            continue
        if row["has_issue"] != 1:
            continue  # skip non-issue rows entirely

        if row["reason_category"]:
            text = render_alasan_ijin(row["reason_category"], row["reason_detail"])
            summary.filled_count += 1
        else:
            text = "NA / Belum ada kabar"
            summary.na_count += 1
        ws.cell(row=r, column=ALASAN_IJIN_COL, value=text)

    out_path = xlsx_path.with_name(f"{xlsx_path.stem} [filled].xlsx")
    wb.save(out_path)
    return out_path, summary

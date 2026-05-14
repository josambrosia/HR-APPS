"""Generate a normalised weekly attendance export (.xlsx) from the DB.

Mirrors the fingerprint structure, cleaned: 12 columns, one row per
attendance record in the date range. Column F (Tipe) shows 'Hari Libur'
for dates marked via the Hari Libur menu (tipe is already stamped in the DB).
"""
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from src.config import DEFAULT_SCHEDULE_START
from src.core.reason_mapper import effective_attendance
from src.db.settings import get_setting, read_lupa_penalty_min

HEADERS = [
    "Nama", "No. Staff", "Dept", "Tanggal", "Hari", "Tipe",
    "Jadwal", "Masuk", "Keluar", "Kerja", "Lembur", "Terlambat",
]
_COL_WIDTHS = [22, 10, 16, 12, 9, 12, 14, 9, 9, 8, 8, 10]


@dataclass
class WeeklyExportSummary:
    rows: int
    employees: int


def generate_weekly_export(
    conn: sqlite3.Connection,
    period_start: str,
    period_end: str,
    out_path: Path,
) -> WeeklyExportSummary:
    """Write a normalised weekly export for [period_start, period_end].

    Holiday rows (tipe='Hari Libur'): column F shows 'Hari Libur', the count
    columns (Masuk/Keluar/Kerja/Lembur/Terlambat) are left blank, Jadwal kept.
    """
    rows = conn.execute(
        """
        SELECT e.nama, e.no_staff, e.dept,
               ar.tanggal, ar.hari, ar.tipe, ar.jadwal,
               ar.masuk, ar.keluar, ar.kerja_jam, ar.lembur_jam,
               ar.terlambat_menit, ar.reason_category
          FROM attendance_records ar
          JOIN employees e ON ar.employee_id = e.id
         WHERE ar.tanggal BETWEEN ? AND ?
         ORDER BY e.nama ASC, ar.tanggal ASC
        """,
        (period_start, period_end),
    ).fetchall()

    schedule_start = get_setting(conn, "schedule_start", default=DEFAULT_SCHEDULE_START)
    lupa_penalty = read_lupa_penalty_min(conn)

    wb = Workbook()
    ws = wb.active
    ws.title = "Mingguan"
    ws.append(HEADERS)
    for col_idx, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)

    employees = set()
    for r in rows:
        employees.add(r["no_staff"])
        is_holiday = r["tipe"] == "Hari Libur"
        if is_holiday:
            eff_masuk, eff_terlambat = None, None
        else:
            eff = effective_attendance(
                r, schedule_start=schedule_start, lupa_penalty_min=lupa_penalty)
            eff_masuk = eff["masuk"]
            eff_terlambat = eff["terlambat_menit"]
        ws.append([
            r["nama"],
            r["no_staff"] or "",
            r["dept"] or "",
            r["tanggal"],
            r["hari"] or "",
            r["tipe"] or "",
            r["jadwal"] or "",
            "" if is_holiday else (eff_masuk or ""),
            "" if is_holiday else (r["keluar"] or ""),
            "" if is_holiday else (
                r["kerja_jam"] if r["kerja_jam"] is not None else ""),
            "" if is_holiday else (
                r["lembur_jam"] if r["lembur_jam"] is not None else ""),
            "" if is_holiday else (
                eff_terlambat if eff_terlambat is not None else ""),
        ])

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return WeeklyExportSummary(rows=len(rows), employees=len(employees))

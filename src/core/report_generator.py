"""Generate Laporan Bulanan {Month} {Year} [Auto Filled].xlsx from DB.

Loads the embedded template (rows 1-2 = headers, rows 3-4 = style samples),
strips the sample rows, then writes data rows + Total Personal per employee
with formatting copied from the samples.
"""
import sqlite3
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime
from typing import Optional

from openpyxl import load_workbook
from openpyxl.cell import MergedCell
from openpyxl.styles import PatternFill

from src.config import TEMPLATE_LAPORAN_BULANAN, REASON_CATEGORIES
from src.core.reason_mapper import render_alasan_ijin


MONTH_NAMES_ID = [
    None, "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]

# Schedule end (16:00) in minutes-from-midnight, for pulang_cepat computation.
JADWAL_END_MINUTES = 16 * 60

# All reason categories except "na" (unknown) and "libur" (a company holiday,
# not a personal ijin) — neither should add to the monthly Ijin column.
IJIN_CATEGORIES = tuple(c for c in REASON_CATEGORIES if c not in ("na", "libur"))

# Gray fill for Total Personal rows — matches reference Laporan Bulanan April.xlsx
# (light gray #C0C0C0 distinguishes total rows from data rows visually).
TOTAL_PERSONAL_FILL = PatternFill(fill_type="solid", fgColor="FFC0C0C0")


@dataclass
class GenerateSummary:
    rows_generated: int
    na_count: int
    employees_count: int


def month_label(year_month: str) -> str:
    """'2026-04' -> 'April 2026'."""
    y, m = year_month.split("-")
    return f"{MONTH_NAMES_ID[int(m)]} {y}"


def hhmm_to_minutes(s: Optional[str]) -> Optional[int]:
    """'08:30' -> 510. Returns None for None/empty/malformed."""
    if not s:
        return None
    try:
        h, m = s.split(":")
        return int(h) * 60 + int(m)
    except (ValueError, AttributeError):
        return None


def compute_derived(row) -> dict:
    """Compute the 5 derived columns from an attendance row (dict-like).

    Returns dict with keys: kurang_jam, pulang_cepat_menit,
    absen_hari, lupa_hari, ijin_hari.
    """
    is_kerja = row.get("tipe") == "Hari Kerja"
    masuk = row.get("masuk")
    keluar = row.get("keluar")
    has_masuk = masuk is not None
    has_keluar = keluar is not None
    kerja_jam = row.get("kerja_jam") or 0
    reason = row.get("reason_category")

    kurang_jam = max(0, 8 - kerja_jam) if is_kerja else 0

    pulang_cepat_menit = 0
    if has_keluar and is_kerja:
        keluar_min = hhmm_to_minutes(keluar)
        if keluar_min is not None and keluar_min < JADWAL_END_MINUTES:
            pulang_cepat_menit = JADWAL_END_MINUTES - keluar_min

    absen_hari = 1 if (is_kerja and not has_masuk and not has_keluar) else 0
    lupa_hari = 1 if (is_kerja and (has_masuk != has_keluar)) else 0
    ijin_hari = 1 if (reason in IJIN_CATEGORIES) else 0

    return {
        "kurang_jam": kurang_jam,
        "pulang_cepat_menit": pulang_cepat_menit,
        "absen_hari": absen_hari,
        "lupa_hari": lupa_hari,
        "ijin_hari": ijin_hari,
    }


def _capture_row_styles(ws, row_num: int, max_col: int = 17) -> list:
    """Return list of style snapshots for cells in `row_num`."""
    styles = []
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        styles.append({
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "number_format": cell.number_format,
        })
    return styles


def _unmerge_row(ws, row_num: int):
    """Remove any merged regions that include `row_num`."""
    to_remove = [
        rng for rng in list(ws.merged_cells.ranges)
        if rng.min_row <= row_num <= rng.max_row
    ]
    for rng in to_remove:
        ws.unmerge_cells(str(rng))


def _apply_row_styles(ws, row_num: int, styles: list):
    """Apply captured styles to cells in `row_num`."""
    for col, style in enumerate(styles, start=1):
        cell = ws.cell(row=row_num, column=col)
        if isinstance(cell, MergedCell):
            continue
        cell.font = style["font"]
        cell.fill = style["fill"]
        cell.border = style["border"]
        cell.alignment = style["alignment"]
        cell.number_format = style["number_format"]


def _write_data_row(ws, row_num: int, db_row, derived: dict, styles: list, *, is_holiday=False):
    """Write 17 cells for one attendance record + apply styles."""
    _unmerge_row(ws, row_num)
    tanggal_val = db_row["tanggal"]
    if isinstance(tanggal_val, str):
        try:
            tanggal_val = datetime.fromisoformat(tanggal_val)
        except ValueError:
            pass

    if is_holiday:
        # Holiday row: marker "Libur" in column G only; Tipe shown as
        # "Hari Kerja" (matches the reference Laporan Bulanan April); all
        # count columns + Alasan Ijin left blank.
        values = [
            db_row["nama"],              # A Nama
            db_row.get("dept") or "",     # B Dept
            tanggal_val,                 # C Tanggal
            db_row.get("hari") or "",     # D Hari
            "Hari Kerja",                # E Tipe (override)
            db_row.get("jadwal") or "",   # F Jadwal
            "Libur",                     # G Masuk -> marker
            "",                          # H Keluar
            "", "", "", "", "", "", "", "",  # I-P counts blank
            "",                          # Q Alasan Ijin blank
        ]
        for col, val in enumerate(values, start=1):
            ws.cell(row=row_num, column=col, value=val)
        _apply_row_styles(ws, row_num, styles)
        return

    has_issue = db_row.get("has_issue") == 1
    reason_cat = db_row.get("reason_category")
    if reason_cat:
        alasan = render_alasan_ijin(reason_cat, db_row.get("reason_detail"))
    elif has_issue:
        alasan = "NA / Belum ada kabar"
    else:
        alasan = None

    # Numeric columns: source fields (kerja_jam, lembur_jam, terlambat_menit) write 0 explicitly;
    # derived counters (absen, lupa, ijin, pulang_cepat, kurang) suppress 0 as visual blank.
    values = [
        db_row["nama"],                                # A: Nama
        db_row.get("dept") or "",                       # B: Dept
        tanggal_val,                                   # C: Tanggal
        db_row.get("hari") or "",                       # D: Hari
        db_row.get("tipe") or "",                       # E: Tipe
        db_row.get("jadwal") or "",                     # F: Jadwal
        db_row.get("masuk") or "",                      # G: Masuk
        db_row.get("keluar") or "",                     # H: Keluar
        db_row.get("kerja_jam") if db_row.get("kerja_jam") is not None else "",
        db_row.get("lembur_jam") if db_row.get("lembur_jam") is not None else "",
        derived["kurang_jam"] if derived["kurang_jam"] else "",
        db_row.get("terlambat_menit") if db_row.get("terlambat_menit") is not None else "",
        derived["pulang_cepat_menit"] if derived["pulang_cepat_menit"] else "",
        derived["absen_hari"] if derived["absen_hari"] else "",
        derived["lupa_hari"] if derived["lupa_hari"] else "",
        derived["ijin_hari"] if derived["ijin_hari"] else "",
        alasan or "",                                   # Q: Alasan Ijin
    ]
    for col, val in enumerate(values, start=1):
        ws.cell(row=row_num, column=col, value=val)
    _apply_row_styles(ws, row_num, styles)


def _init_total() -> dict:
    return {
        "kerja_jam": 0.0, "lembur_jam": 0.0, "kurang_jam": 0.0,
        "terlambat_menit": 0, "pulang_cepat_menit": 0,
        "absen_hari": 0, "lupa_hari": 0, "ijin_hari": 0,
    }


def _accumulate_total(total: dict, db_row, derived: dict):
    total["kerja_jam"] += db_row.get("kerja_jam") or 0
    total["lembur_jam"] += db_row.get("lembur_jam") or 0
    total["kurang_jam"] += derived["kurang_jam"]
    total["terlambat_menit"] += db_row.get("terlambat_menit") or 0
    total["pulang_cepat_menit"] += derived["pulang_cepat_menit"]
    total["absen_hari"] += derived["absen_hari"]
    total["lupa_hari"] += derived["lupa_hari"]
    total["ijin_hari"] += derived["ijin_hari"]


def _write_total_row(ws, row_num: int, total: dict, styles: list):
    """Write Total Personal row. Cell A holds the label.

    Total row layout from template row 4: A spans A:H ("Total Personal:"),
    then I-P have the sums. Q is blank.
    """
    _unmerge_row(ws, row_num)
    # Columns B-H (2-8) intentionally left blank — they get merged with A
    # below to form the "Total Personal:" label span.
    ws.cell(row=row_num, column=1, value="Total Personal:")
    ws.cell(row=row_num, column=9, value=round(total["kerja_jam"], 1))
    ws.cell(row=row_num, column=10, value=round(total["lembur_jam"], 1))
    ws.cell(row=row_num, column=11, value=round(total["kurang_jam"], 1))
    ws.cell(row=row_num, column=12, value=total["terlambat_menit"])
    ws.cell(row=row_num, column=13, value=total["pulang_cepat_menit"])
    ws.cell(row=row_num, column=14, value=total["absen_hari"])
    ws.cell(row=row_num, column=15, value=total["lupa_hari"])
    ws.cell(row=row_num, column=16, value=total["ijin_hari"])
    _apply_row_styles(ws, row_num, styles)
    # Override fill to gray — template's row 4 sample is a data row (white),
    # but the original Laporan Bulanan uses light gray for Total Personal rows.
    for col in range(1, 18):
        cell = ws.cell(row=row_num, column=col)
        if not isinstance(cell, MergedCell):
            cell.fill = TOTAL_PERSONAL_FILL
    # Re-merge A:H to match template's Total Personal layout
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=8)


def generate_monthly_report(
    conn: sqlite3.Connection,
    year_month: str,
    out_path: Path,
) -> GenerateSummary:
    """Generate Laporan Bulanan for the given year-month from DB.

    Returns GenerateSummary with rows_generated, na_count, employees_count.
    """
    wb = load_workbook(TEMPLATE_LAPORAN_BULANAN)
    ws = wb.active

    # Capture styles from sample rows 3 (data) and 4 (Total Personal)
    data_styles = _capture_row_styles(ws, 3)
    total_styles = _capture_row_styles(ws, 4)

    # Remove all merged regions that fall below the header (row > 2)
    # BEFORE deleting rows, to avoid openpyxl's internal shift-bug leaving
    # phantom MergedCell objects in the cell cache.
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row > 2:
            ws.unmerge_cells(str(rng))

    # Strip the sample rows (rows 3 and 4)
    ws.delete_rows(3, 2)

    # Query data
    cursor = conn.execute(
        """
        SELECT ar.*, e.nama, e.dept
          FROM attendance_records ar
          JOIN employees e ON ar.employee_id = e.id
         WHERE substr(ar.tanggal, 1, 7) = ?
         ORDER BY e.nama, ar.tanggal
        """,
        (year_month,),
    )
    rows = [dict(r) for r in cursor.fetchall()]

    # Group by employee (preserving order)
    employee_blocks = []
    current_emp = None
    current_block = []
    for r in rows:
        if r["nama"] != current_emp:
            if current_block:
                employee_blocks.append((current_emp, current_block))
            current_emp = r["nama"]
            current_block = []
        current_block.append(r)
    if current_block:
        employee_blocks.append((current_emp, current_block))

    # Write rows
    current_row = 3
    rows_generated = 0
    na_count = 0

    for _emp_name, emp_records in employee_blocks:
        total = _init_total()
        for r in emp_records:
            is_holiday = r.get("tipe") == "Hari Libur"
            derived = {} if is_holiday else compute_derived(r)
            _write_data_row(ws, current_row, r, derived, data_styles, is_holiday=is_holiday)
            if not is_holiday:
                _accumulate_total(total, r, derived)
                if r.get("has_issue") == 1 and not r.get("reason_category"):
                    na_count += 1
            rows_generated += 1
            current_row += 1
        _write_total_row(ws, current_row, total, total_styles)
        current_row += 1

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    return GenerateSummary(
        rows_generated=rows_generated,
        na_count=na_count,
        employees_count=len(employee_blocks),
    )

"""v26.0.2 — aplikasi menghitung sendiri telat & lembur dari jam absensi.

Cakupan:
- Task 3: pemecahan riwayat jadwal historis (DB instalasi baru berisi
  data lama, atau import membawa tanggal lama).
- Task 2: recompute sekali saat upgrade (flag idempoten, skip editan manual).
- Task 1: recompute saat import (kolom file diabaikan).
- Task 4: kolom Jadwal di kedua export = jadwal aplikasi (read-time).
- Task 6: pembanding db vs app di Ekspor Diagnostik.
"""
import sqlite3

from openpyxl import load_workbook

from src.core.attendance_calc import recompute
from src.core.diagnostics import _samples_rows
from src.core.report_generator import generate_monthly_report
from src.core.weekly_export import generate_weekly_export
from src.db.attendance import upsert_attendance
from src.db.connection import get_connection
from src.db.employees import upsert_employee
from src.db.recompute import FLAG_KEY
from src.db.schedule import ensure_historical_schedule, list_periods, schedule_for
from src.db.schema import init_db
from src.parsers.fingerprint import FingerprintRow
from src.ui.screens.import_screen import run_import_confirm


def _legacy_db(p, *, sched=("08.00", "16.00")):
    with sqlite3.connect(p) as conn:
        conn.execute("CREATE TABLE settings (key TEXT PRIMARY KEY, value TEXT)")
        conn.execute("INSERT INTO settings VALUES ('schedule_start',?)",
                     (sched[0],))
        conn.execute("INSERT INTO settings VALUES ('schedule_end',?)",
                     (sched[1],))
        conn.commit()
    init_db(p)


def _att(conn, eid, tanggal, masuk, keluar, *, terlambat=None, lembur=None,
         kerja=None, manual=False):
    upsert_attendance(
        conn, employee_id=eid, tanggal=tanggal, hari=None, tipe="Hari Kerja",
        jadwal="08.00 - 16.00", masuk=masuk, keluar=keluar,
        kerja_jam=kerja, lembur_jam=lembur, terlambat_menit=terlambat,
        has_issue=0, imported_from="t")
    if manual:
        conn.execute(
            "UPDATE attendance_records SET manual_edited_at = '2026-01-01' "
            "WHERE employee_id = ? AND tanggal = ?", (eid, tanggal))


def _drop_flag(p):
    with sqlite3.connect(p) as conn:
        conn.execute("DELETE FROM settings WHERE key = ?", (FLAG_KEY,))
        conn.commit()


def _periods(p):
    with get_connection(p) as conn:
        return [(r["effective_from"], r["start_time"], r["end_time"])
                for r in list_periods(conn)]


def _fp(no_staff, tanggal, **kw):
    base = dict(nama="Budi", no_staff=no_staff, dept="Prod", tanggal=tanggal,
                hari="Rabu", tipe="Hari Kerja", jadwal="08.00 - 16.00",
                masuk="08:00", keluar="16:00", kerja_jam=8.0, lembur_jam=0.0,
                terlambat_menit=0, source_file="f.xls")
    base.update(kw)
    return FingerprintRow(**base)


# ── Task 3: riwayat jadwal historis ────────────────────────────────────────

def test_ensure_historical_schedule_splits_default_period_with_old_rows(
        tmp_path):
    """Instalasi baru (1 periode default) + baris < cutoff → dipecah."""
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        _att(conn, eid, "2026-09-10", "08:14", "16:30")
    assert _periods(p) == [("1970-01-01", "09.00", "17.00")]
    init_db(p)  # app start berikutnya
    assert _periods(p) == [
        ("1970-01-01", "08.00", "16.00"),
        ("2026-09-14", "09.00", "17.00"),
    ]


def test_ensure_historical_schedule_noop_without_old_rows(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        _att(conn, eid, "2026-09-15", "09:05", "17:30")
    init_db(p)
    assert _periods(p) == [("1970-01-01", "09.00", "17.00")]


def test_ensure_historical_schedule_noop_when_already_split(tmp_path):
    p = tmp_path / "hr.db"
    _legacy_db(p)  # settings lama → 2 periode sejak seeding
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        _att(conn, eid, "2026-09-10", "08:14", "16:30")
    init_db(p)
    assert _periods(p) == [
        ("1970-01-01", "08.00", "16.00"),
        ("2026-09-14", "09.00", "17.00"),
    ]


def test_ensure_historical_schedule_noop_for_custom_schedule(tmp_path):
    p = tmp_path / "hr.db"
    _legacy_db(p, sched=("08.30", "16.30"))
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        _att(conn, eid, "2026-09-10", "08:40", "16:40")
    init_db(p)
    assert _periods(p) == [("1970-01-01", "08.30", "16.30")]


def test_ensure_historical_schedule_extra_dates_split(tmp_path):
    """Import membawa tanggal lama ke DB instalasi baru → pecah."""
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        changed = ensure_historical_schedule(
            conn, extra_dates={"2026-05-06", "2026-05-07"})
    assert changed is True
    assert _periods(p) == [
        ("1970-01-01", "08.00", "16.00"),
        ("2026-09-14", "09.00", "17.00"),
    ]


# ── Task 2: recompute sekali saat upgrade ──────────────────────────────────

def test_upgrade_recomputes_with_per_date_schedule(tmp_path):
    """DB instalasi baru berisi data lama: upgrade memecah jadwal +
    menghitung ulang telat dari jam absensi (basis lama vs baru)."""
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        # Nilai tersimpan = angka mesin (masuk − 08:00)
        _att(conn, eid, "2026-09-10", "08:14", "16:30", terlambat=14, lembur=0.5)
        _att(conn, eid, "2026-09-15", "08:14", "17:30", terlambat=14, lembur=1.5)
    _drop_flag(p)  # simulasikan DB yang belum pernah di-recompute v26.0.2
    init_db(p)
    with get_connection(p) as conn:
        r10 = conn.execute(
            "SELECT terlambat_menit, lembur_jam FROM attendance_records "
            "WHERE tanggal='2026-09-10'").fetchone()
        r15 = conn.execute(
            "SELECT terlambat_menit, lembur_jam FROM attendance_records "
            "WHERE tanggal='2026-09-15'").fetchone()
    assert (r10["terlambat_menit"], r10["lembur_jam"]) == (14, 0.5)
    # 15 Sep basis 09.00: 08:14 tidak telat, lembur 17:30-17:00 = 0.5
    assert (r15["terlambat_menit"], r15["lembur_jam"]) == (0, 0.5)


def test_recompute_skips_manual_rows(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        _att(conn, eid, "2026-09-15", "08:14", "17:30",
             terlambat=7, lembur=7.0, manual=True)
    _drop_flag(p)
    init_db(p)
    with get_connection(p) as conn:
        r = conn.execute(
            "SELECT terlambat_menit, lembur_jam FROM attendance_records"
        ).fetchone()
    assert (r["terlambat_menit"], r["lembur_jam"]) == (7, 7.0)


def test_recompute_does_not_touch_kerja_jam_and_has_issue(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        _att(conn, eid, "2026-09-15", "09:05", "17:30",
             terlambat=61, lembur=0.5, kerja=8.3)
    _drop_flag(p)
    init_db(p)
    with get_connection(p) as conn:
        r = conn.execute(
            "SELECT kerja_jam, has_issue, terlambat_menit FROM attendance_records"
        ).fetchone()
    assert r["kerja_jam"] == 8.3
    assert r["has_issue"] == 0
    assert r["terlambat_menit"] == 5


def test_recompute_flag_idempotent(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        _att(conn, eid, "2026-09-15", "09:05", "17:30", terlambat=61)
    _drop_flag(p)
    init_db(p)   # recompute jalan, flag dipasang
    with get_connection(p) as conn:
        assert conn.execute(
            "SELECT 1 FROM settings WHERE key = ?", (FLAG_KEY,)
        ).fetchone() is not None
        conn.execute(
            "UPDATE attendance_records SET terlambat_menit = 999")
    init_db(p)   # flag ada → tidak menghitung ulang
    with get_connection(p) as conn:
        r = conn.execute(
            "SELECT terlambat_menit FROM attendance_records").fetchone()
    assert r["terlambat_menit"] == 999


def test_fresh_db_flag_present_after_init(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    with sqlite3.connect(p) as conn:
        row = conn.execute(
            "SELECT 1 FROM settings WHERE key = ?", (FLAG_KEY,)).fetchone()
    assert row is not None


# ── Task 1: recompute saat import ──────────────────────────────────────────

def test_import_ignores_file_terlambat_and_lembur(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    rows = [_fp("1", "2026-09-15", masuk="09:05", keluar="19:30",
                kerja_jam=8.0, lembur_jam=1.0, terlambat_menit=99)]
    run_import_confirm(p, rows)
    with get_connection(p) as conn:
        r = conn.execute(
            "SELECT terlambat_menit, lembur_jam, kerja_jam "
            "FROM attendance_records").fetchone()
    assert r["terlambat_menit"] == 5     # 09:05 − 09:00
    assert r["lembur_jam"] == 2.5        # 19:30 − 17:00 (file bilang 1.0)
    assert r["kerja_jam"] == 8.0         # tetap dari file


def test_import_old_dates_split_schedule_and_use_old_basis(tmp_path):
    """Import tanggal < cutoff ke instalasi baru: jadwal dipecah, telat
    dihitung dari basis mesin lama 08.00."""
    p = tmp_path / "hr.db"
    init_db(p)
    rows = [_fp("1", "2026-09-10", masuk="08:14", keluar="16:30",
                lembur_jam=0.0, terlambat_menit=0)]
    run_import_confirm(p, rows)
    with get_connection(p) as conn:
        r = conn.execute(
            "SELECT terlambat_menit, lembur_jam FROM attendance_records"
        ).fetchone()
        sched = schedule_for(conn, "2026-09-10")
    assert sched == ("08.00", "16.00")
    assert r["terlambat_menit"] == 14    # 08:14 − 08:00
    assert r["lembur_jam"] == 0.5        # 16:30 − 16:00


def test_import_mixed_dates_per_date_basis(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    rows = [
        _fp("1", "2026-09-10", masuk="08:14", keluar="16:30"),
        _fp("1", "2026-09-15", masuk="08:14", keluar="16:30"),
    ]
    run_import_confirm(p, rows)
    with get_connection(p) as conn:
        old = conn.execute(
            "SELECT terlambat_menit FROM attendance_records "
            "WHERE tanggal='2026-09-10'").fetchone()
        new = conn.execute(
            "SELECT terlambat_menit FROM attendance_records "
            "WHERE tanggal='2026-09-15'").fetchone()
    assert old["terlambat_menit"] == 14   # basis 08.00
    assert new["terlambat_menit"] == 0    # basis 09.00


# ── Task 4: kolom Jadwal read-time di export ───────────────────────────────

def test_monthly_report_jadwal_column_uses_app_schedule(tmp_path):
    p = tmp_path / "hr.db"
    _legacy_db(p)  # 2 periode: 08-16 lalu 09-17
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="ANDI")
        _att(conn, eid, "2026-09-10", "08:00", "16:30")
        _att(conn, eid, "2026-09-15", "09:00", "16:30")
        out = tmp_path / "out.xlsx"
        generate_monthly_report(conn, year_month="2026-09", out_path=out)
    ws = load_workbook(out).active
    assert ws.cell(row=3, column=6).value == "08.00 - 16.00"   # F Jadwal
    assert ws.cell(row=4, column=6).value == "09.00 - 17.00"


def test_weekly_export_jadwal_column_uses_app_schedule(tmp_path):
    p = tmp_path / "hr.db"
    _legacy_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="ANDI")
        _att(conn, eid, "2026-09-10", "08:00", "16:30")
        _att(conn, eid, "2026-09-15", "09:00", "16:30")
        out = tmp_path / "weekly.xlsx"
        generate_weekly_export(conn, "2026-09-01", "2026-09-30", out)
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=7).value == "08.00 - 16.00"   # G Jadwal
    assert ws.cell(row=3, column=7).value == "09.00 - 17.00"


# ── Task 6: pembanding db vs app di diagnostics ────────────────────────────

def test_diagnostics_samples_compare_db_vs_app(tmp_path):
    p = tmp_path / "hr.db"
    init_db(p)
    with get_connection(p) as conn:
        eid = upsert_employee(conn, no_staff="1", nama="Agus")
        # Nilai DB menyimpan angka mesin lama; app menghitung 5 menit.
        _att(conn, eid, "2026-09-15", "09:05", "18:30",
             terlambat=65, lembur=2.5)
        rows = _samples_rows(conn)
    assert len(rows) == 1
    r = rows[0]
    assert r["terlambat_db"] == 65
    assert r["terlambat_app"] == 5       # 09:05 − 09:00
    assert r["lembur_db"] == 2.5
    assert r["lembur_app"] == 1.5        # 18:30 − 17:00


def test_diagnostics_app_matches_recompute_helper(tmp_path):
    """Nilai app di CSV = hasil attendance_calc.recompute yang sama."""
    calc = recompute(tipe="Hari Kerja", masuk="09:05", keluar="18:30",
                     schedule_start="09.00", schedule_end="17.00")
    assert calc["terlambat_menit"] == 5
    assert calc["lembur_jam"] == 1.5

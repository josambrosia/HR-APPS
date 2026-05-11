"""Tests for report_generator: derived fields, totals, end-to-end."""
import sqlite3
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.db.schema import DDL
from src.db.employees import upsert_employee
from src.db.attendance import upsert_attendance, set_reason
from src.core.report_generator import (
    compute_derived, hhmm_to_minutes,
    generate_monthly_report, month_label, MONTH_NAMES_ID,
    GenerateSummary,
)


# ────────────────────────────────────────────── Pure helpers


def test_hhmm_to_minutes_basic():
    assert hhmm_to_minutes("08:00") == 480
    assert hhmm_to_minutes("16:30") == 990
    assert hhmm_to_minutes("00:00") == 0


def test_hhmm_to_minutes_none_returns_none():
    assert hhmm_to_minutes(None) is None
    assert hhmm_to_minutes("") is None


def test_month_label_indonesian():
    assert month_label("2026-04") == "April 2026"
    assert month_label("2026-12") == "Desember 2026"
    assert month_label("2027-01") == "Januari 2027"


# ────────────────────────────────────────────── compute_derived


def _row(tipe="Hari Kerja", masuk="08:00", keluar="16:00",
         kerja_jam=8.0, terlambat_menit=0, reason_category=None):
    return {
        "tipe": tipe, "masuk": masuk, "keluar": keluar,
        "kerja_jam": kerja_jam, "terlambat_menit": terlambat_menit,
        "reason_category": reason_category,
    }


def test_compute_derived_full_workday():
    d = compute_derived(_row())
    assert d == {"kurang_jam": 0, "pulang_cepat_menit": 0,
                 "absen_hari": 0, "lupa_hari": 0, "ijin_hari": 0}


def test_compute_derived_absen_full_day():
    d = compute_derived(_row(masuk=None, keluar=None, kerja_jam=None))
    assert d["absen_hari"] == 1
    assert d["lupa_hari"] == 0
    assert d["kurang_jam"] == 8


def test_compute_derived_lupa_in_or_out():
    # Only masuk, no keluar
    d = compute_derived(_row(keluar=None, kerja_jam=4.0))
    assert d["lupa_hari"] == 1
    assert d["absen_hari"] == 0
    assert d["kurang_jam"] == 4


def test_compute_derived_pulang_cepat():
    d = compute_derived(_row(keluar="15:30", kerja_jam=7.5))
    assert d["pulang_cepat_menit"] == 30
    assert abs(d["kurang_jam"] - 0.5) < 1e-9


def test_compute_derived_istirahat_no_derived():
    d = compute_derived(_row(tipe="Istirahat", masuk=None, keluar=None, kerja_jam=None))
    assert d["absen_hari"] == 0   # istirahat day, not Hari Kerja
    assert d["kurang_jam"] == 0
    assert d["lupa_hari"] == 0


def test_compute_derived_ijin_marker():
    d = compute_derived(_row(masuk=None, keluar=None, reason_category="izin_sakit"))
    assert d["ijin_hari"] == 1
    d2 = compute_derived(_row(masuk=None, keluar=None, reason_category=None))
    assert d2["ijin_hari"] == 0


# ────────────────────────────────────────────── End-to-end generate


def _conn():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.executescript(DDL)
    return conn


def _populate_april(conn):
    """Add 1 employee with 3 days of attendance in April 2026."""
    emp_id = upsert_employee(conn, no_staff="E001", nama="ANDIKA",
                              dept="ARGA DIRGA", phone=None)
    # Day 1: normal day
    upsert_attendance(conn, employee_id=emp_id, tanggal="2026-04-01",
                      hari="Rabu", tipe="Hari Kerja", jadwal="08.00 - 16.00",
                      masuk="08:06", keluar="16:30",
                      kerja_jam=7.9, lembur_jam=0.5, terlambat_menit=6,
                      has_issue=0, imported_from="test.xls")
    # Day 2: absen (masuk null, keluar null)
    upsert_attendance(conn, employee_id=emp_id, tanggal="2026-04-02",
                      hari="Kamis", tipe="Hari Kerja", jadwal="08.00 - 16.00",
                      masuk=None, keluar=None,
                      kerja_jam=None, lembur_jam=None, terlambat_menit=None,
                      has_issue=1, imported_from="test.xls")
    # Day 3: resolved issue with izin_sakit
    upsert_attendance(conn, employee_id=emp_id, tanggal="2026-04-03",
                      hari="Jumat", tipe="Hari Kerja", jadwal="08.00 - 16.00",
                      masuk=None, keluar=None,
                      kerja_jam=None, lembur_jam=None, terlambat_menit=None,
                      has_issue=1, imported_from="test.xls")
    row_id = conn.execute(
        "SELECT id FROM attendance_records WHERE tanggal='2026-04-03'"
    ).fetchone()["id"]
    set_reason(conn, attendance_id=row_id, category="izin_sakit", detail=None)
    return emp_id


def test_generate_summary_counts():
    conn = _conn()
    _populate_april(conn)
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "test.xlsx"
        summary = generate_monthly_report(conn, year_month="2026-04", out_path=out)
        assert isinstance(summary, GenerateSummary)
        assert summary.rows_generated == 3
        assert summary.na_count == 1   # day 2 still open
        assert summary.employees_count == 1
        assert out.exists()


def test_generate_creates_correct_file_structure():
    conn = _conn()
    _populate_april(conn)
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "test.xlsx"
        generate_monthly_report(conn, year_month="2026-04", out_path=out)
        wb = load_workbook(out)
        ws = wb.active

        # Row 1: header
        assert ws.cell(row=1, column=1).value == "Nama"
        assert ws.cell(row=1, column=17).value == "Alasan Ijin"
        # Row 3: first data row (ANDIKA, 2026-04-01)
        assert ws.cell(row=3, column=1).value == "ANDIKA"
        assert ws.cell(row=3, column=2).value == "ARGA DIRGA"
        # Row 5: third data row (2026-04-03) — has reason Izin Sakit
        assert "Sakit" in str(ws.cell(row=5, column=17).value)
        # Total Personal row exists somewhere after the data rows
        found_total = False
        for r in range(3, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == "Total Personal:":
                found_total = True
                break
        assert found_total, "Total Personal row not found"


def test_generate_empty_month_writes_header_only():
    conn = _conn()
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "empty.xlsx"
        summary = generate_monthly_report(conn, year_month="2026-04", out_path=out)
        assert summary.rows_generated == 0
        assert summary.employees_count == 0
        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Nama"
        # No data rows beyond header (row 3 should be empty)
        assert ws.cell(row=3, column=1).value in (None, "")


def test_generate_na_for_open_issues_in_alasan_column():
    conn = _conn()
    _populate_april(conn)
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "test.xlsx"
        generate_monthly_report(conn, year_month="2026-04", out_path=out)
        wb = load_workbook(out)
        ws = wb.active
        # Day 2 (Apr 2) is open issue → should have "NA / Belum ada kabar"
        from datetime import datetime
        for r in range(3, ws.max_row + 1):
            tanggal = ws.cell(row=r, column=3).value
            # tanggal could be datetime, date, or string
            day_2_match = False
            if isinstance(tanggal, datetime) and tanggal.day == 2 and tanggal.month == 4:
                day_2_match = True
            elif tanggal and str(tanggal).startswith("2026-04-02"):
                day_2_match = True
            if day_2_match:
                assert "NA" in str(ws.cell(row=r, column=17).value)
                return
        pytest.fail("Apr 2 row not found in generated file")

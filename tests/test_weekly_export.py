"""Tests for src/core/weekly_export.py — normalised weekly export."""
import sqlite3

from openpyxl import load_workbook

from src.db.schema import DDL
from src.db.employees import upsert_employee
from src.db.attendance import upsert_attendance
from src.db.holidays import mark_holidays
from src.core.weekly_export import generate_weekly_export, HEADERS


def _conn():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.executescript(DDL)
    return conn


def _att(conn, emp_id, tanggal, hari):
    upsert_attendance(
        conn, employee_id=emp_id, tanggal=tanggal, hari=hari,
        tipe="Hari Kerja", jadwal="08.00 - 16.00",
        masuk="08:05", keluar="16:00", kerja_jam=8.0,
        lembur_jam=None, terlambat_menit=5, has_issue=0,
        imported_from="W1.xls",
    )


def test_weekly_export_has_header_and_12_columns(tmp_path):
    conn = _conn()
    a = upsert_employee(conn, no_staff="1", nama="ANDI", dept="X")
    _att(conn, a, "2026-04-07", "Selasa")
    out = tmp_path / "weekly.xlsx"
    generate_weekly_export(conn, "2026-04-06", "2026-04-12", out)
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == HEADERS
    assert len(HEADERS) == 12


def test_weekly_export_filters_date_range(tmp_path):
    conn = _conn()
    a = upsert_employee(conn, no_staff="1", nama="ANDI", dept="X")
    _att(conn, a, "2026-04-07", "Selasa")   # inside
    _att(conn, a, "2026-04-20", "Senin")    # outside
    out = tmp_path / "weekly.xlsx"
    summary = generate_weekly_export(conn, "2026-04-06", "2026-04-12", out)
    assert summary.rows == 1
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=4).value == "2026-04-07"   # D Tanggal


def test_weekly_export_holiday_row_tipe_and_blank_counts(tmp_path):
    conn = _conn()
    a = upsert_employee(conn, no_staff="1", nama="ANDI", dept="X")
    _att(conn, a, "2026-04-09", "Kamis")
    mark_holidays(conn, ["2026-04-09"])
    out = tmp_path / "weekly.xlsx"
    generate_weekly_export(conn, "2026-04-06", "2026-04-12", out)
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=6).value == "Hari Libur"      # F Tipe
    assert ws.cell(row=2, column=7).value == "08.00 - 16.00"   # G Jadwal kept
    assert ws.cell(row=2, column=8).value in (None, "")        # H Masuk blank
    assert ws.cell(row=2, column=12).value in (None, "")       # L Terlambat blank


def test_weekly_export_sorted_by_nama_then_tanggal(tmp_path):
    conn = _conn()
    a = upsert_employee(conn, no_staff="1", nama="ZARA", dept="X")
    b = upsert_employee(conn, no_staff="2", nama="ANDI", dept="X")
    _att(conn, a, "2026-04-07", "Selasa")
    _att(conn, b, "2026-04-08", "Rabu")
    _att(conn, b, "2026-04-07", "Selasa")
    out = tmp_path / "weekly.xlsx"
    generate_weekly_export(conn, "2026-04-06", "2026-04-12", out)
    ws = load_workbook(out).active
    assert [ws.cell(row=r, column=1).value for r in range(2, 5)] == \
        ["ANDI", "ANDI", "ZARA"]


def test_weekly_export_summary_counts_distinct_staff_not_names(tmp_path):
    """Two employees sharing a name must count as 2, not 1."""
    conn = _conn()
    a = upsert_employee(conn, no_staff="1", nama="BUDI", dept="X")
    b = upsert_employee(conn, no_staff="2", nama="BUDI", dept="Y")
    _att(conn, a, "2026-04-07", "Selasa")
    _att(conn, b, "2026-04-07", "Selasa")
    out = tmp_path / "weekly.xlsx"
    summary = generate_weekly_export(conn, "2026-04-06", "2026-04-12", out)
    assert summary.rows == 2
    assert summary.employees == 2   # would be 1 if deduped by name


def test_weekly_export_effective_masuk_work_justified_late(tmp_path):
    from src.db.attendance import set_reason
    conn = _conn()
    a = upsert_employee(conn, no_staff="1", nama="ANDI", dept="X")
    upsert_attendance(
        conn, employee_id=a, tanggal="2026-04-07", hari="Selasa",
        tipe="Hari Kerja", jadwal="08.00 - 16.00", masuk="09:40",
        keluar="16:00", kerja_jam=6.5, lembur_jam=None,
        terlambat_menit=100, has_issue=1, imported_from="W1.xls",
    )
    rid = conn.execute("SELECT id FROM attendance_records").fetchone()["id"]
    set_reason(conn, attendance_id=rid, category="tugas_paparan", detail="PT X")
    out = tmp_path / "weekly.xlsx"
    generate_weekly_export(conn, "2026-04-06", "2026-04-12", out)
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=8).value == "08:00"   # H Masuk -> effective
    assert ws.cell(row=2, column=12).value == 0        # L Terlambat -> 0


def test_weekly_export_effective_forgot_clock_in(tmp_path):
    from src.db.attendance import set_reason
    conn = _conn()
    a = upsert_employee(conn, no_staff="1", nama="ANDI", dept="X")
    upsert_attendance(
        conn, employee_id=a, tanggal="2026-04-07", hari="Selasa",
        tipe="Hari Kerja", jadwal="08.00 - 16.00", masuk=None,
        keluar="16:05", kerja_jam=None, lembur_jam=None,
        terlambat_menit=None, has_issue=1, imported_from="W1.xls",
    )
    rid = conn.execute("SELECT id FROM attendance_records").fetchone()["id"]
    set_reason(conn, attendance_id=rid, category="lupa_absen_datang", detail=None)
    out = tmp_path / "weekly.xlsx"
    generate_weekly_export(conn, "2026-04-06", "2026-04-12", out)
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=8).value == "08:15"   # H Masuk -> effective
    assert ws.cell(row=2, column=12).value == 15       # L Terlambat -> penalty

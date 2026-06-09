# tests/test_report_filler_severe_lateness.py
from pathlib import Path
from openpyxl import Workbook
from src.db.schema import init_db
from src.db.connection import get_connection
from src.db.employees import upsert_employee
from src.db.attendance import upsert_attendance, set_reason
from src.core.report_filler import fill_monthly_report, ALASAN_IJIN_COL


def _make_template(tmp_path, nama, tanggal_str):
    wb = Workbook()
    ws = wb.active
    # rows 1-2 are headers; data starts row 3. Col 1 = nama, col 3 = tanggal.
    ws.cell(row=3, column=1, value=nama)
    ws.cell(row=3, column=3, value=tanggal_str)
    p = tmp_path / "tpl.xlsx"
    wb.save(p)
    return p


def test_resolved_severe_lateness_writes_alasan(tmp_path, temp_db_path):
    init_db(temp_db_path)
    with get_connection(temp_db_path) as conn:
        a = upsert_employee(conn, no_staff="1", nama="BUDI", dept="IT")
        upsert_attendance(conn, employee_id=a, tanggal="2026-05-04", hari="Senin",
                          tipe="Hari Kerja", jadwal="08.00 - 16.00",
                          masuk="09:15", keluar="16:30", kerja_jam=None,
                          lembur_jam=None, terlambat_menit=75, has_issue=0,
                          imported_from="W1.xls")
        rid = conn.execute("SELECT id FROM attendance_records").fetchone()["id"]
        set_reason(conn, attendance_id=rid, category="terlambat_lain", detail="macet")
        tpl = _make_template(tmp_path, "BUDI", "04/05/2026")
        out, summary = fill_monthly_report(tpl, conn, out_dir=tmp_path)
    from openpyxl import load_workbook
    ws = load_workbook(out).active
    val = ws.cell(row=3, column=ALASAN_IJIN_COL).value
    assert val and "Terlambat" in val
    assert summary.filled_count == 1


def test_unresolved_issue_still_writes_na(tmp_path, temp_db_path):
    init_db(temp_db_path)
    with get_connection(temp_db_path) as conn:
        a = upsert_employee(conn, no_staff="1", nama="ANI", dept="HR")
        upsert_attendance(conn, employee_id=a, tanggal="2026-05-05", hari="Selasa",
                          tipe="Hari Kerja", jadwal="08.00 - 16.00",
                          masuk=None, keluar="16:00", kerja_jam=None,
                          lembur_jam=None, terlambat_menit=None, has_issue=1,
                          imported_from="W1.xls")
        tpl = _make_template(tmp_path, "ANI", "05/05/2026")
        out, summary = fill_monthly_report(tpl, conn, out_dir=tmp_path)
    from openpyxl import load_workbook
    ws = load_workbook(out).active
    assert ws.cell(row=3, column=ALASAN_IJIN_COL).value == "NA / Belum ada kabar"
    assert summary.na_count == 1


def test_unresolved_non_issue_row_skipped(tmp_path, temp_db_path):
    init_db(temp_db_path)
    with get_connection(temp_db_path) as conn:
        a = upsert_employee(conn, no_staff="1", nama="CITRA", dept="Fin")
        upsert_attendance(conn, employee_id=a, tanggal="2026-05-06", hari="Rabu",
                          tipe="Hari Kerja", jadwal="08.00 - 16.00",
                          masuk="09:10", keluar="17:00", kerja_jam=None,
                          lembur_jam=None, terlambat_menit=70, has_issue=0,
                          imported_from="W1.xls")
        tpl = _make_template(tmp_path, "CITRA", "06/05/2026")
        out, summary = fill_monthly_report(tpl, conn, out_dir=tmp_path)
    from openpyxl import load_workbook
    ws = load_workbook(out).active
    assert ws.cell(row=3, column=ALASAN_IJIN_COL).value in (None, "")
    assert summary.filled_count == 0 and summary.na_count == 0
